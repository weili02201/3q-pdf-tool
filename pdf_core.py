# -*- coding: utf-8 -*-
# SPDX-License-Identifier: AGPL-3.0-only
# Copyright (C) 2026 三扣科技 (3Q Tech)
# 本作品基于 PyMuPDF（Copyright (C) Artifex Software, Inc.，AGPL-3.0 或商业许可），
# 依 AGPL-3.0 第 5 条整体以 AGPL-3.0 授权。完整文本见同目录 LICENSE，第三方声明见 NOTICE。
"""
3Q PDF开源免费工具 · 核心层（pdf_core）

这一层是**人和 Agent 共用的唯一事实源**：
- 只依赖 fitz(PyMuPDF) / Pillow / openpyxl / python-docx，完全不 import tkinter
- 因此能在无图形界面的环境（服务器、CI、Agent 子线程）直接 import
- 人类入口 pdf_tool.py 与 Agent 入口 pdf_cli.py 都 import 这一层，逻辑不会漂移成两份

改造方式（2026-10-18）：从 pdf_tool.py 按行区间切出，内容是原样搬运，未改写逻辑。
"""
import io
import json
import math
import os
import re
import sys
import tempfile
import time

import fitz  # PyMuPDF


VERSION = "V2.0"
MM = 72 / 25.4  # 毫米 -> PDF 点
FONT_CJK = "china-ss"  # PyMuPDF 内置中文字体（黑体风格），无需系统字体

COLOR_PRESETS = {
    "浅灰": (0.75, 0.75, 0.75),
    "黑": (0, 0, 0),
    "红": (0.8, 0.1, 0.1),
    "蓝": (0.1, 0.3, 0.8),
}

# 画笔预设色（界面色块与 CLI --color 共用一份，别再各写一套）
INK_PRESETS = {
    "红": "#E11D48",
    "黑": "#1F2937",
    "蓝": "#2563EB",
    "黄": "#F59E0B",
    "绿": "#0E9F6E",
}

# 画笔形状（apply_ink 的 shape 取值，界面标签共用一份）
INK_SHAPES = ("free", "line", "rect", "ellipse")
INK_SHAPE_LABELS = {"free": "手绘", "line": "直线",
                    "rect": "矩形", "ellipse": "椭圆"}


def hex_to_rgb(h):
    """"#RRGGBB" 或 "RRGGBB" → (r,g,b) 0..1 三元组。供画笔/CLI 共用。"""
    if isinstance(h, (tuple, list)) and len(h) == 3:
        return tuple(float(c) for c in h)
    s = str(h or "").strip().lstrip("#")
    if len(s) != 6:
        raise ValueError("颜色应为 #RRGGBB 六位十六进制，如 #E11D48")
    return tuple(int(s[i:i + 2], 16) / 255.0 for i in (0, 2, 4))

# 文字以「图片」方式贴入页面：预览与输出共用同一张光栅图，
# 位置即图片左上角，彻底消除字体基线/降部带来的偏差，做到所见即所得。
TEXT_SS = 4  # 文字光栅超采样倍数（1 图像像素 = 1/TEXT_SS 点）
_TEXT_CACHE = {}


def _cjk_font_bytes():
    """内置中文字体字节（PyMuPDF china-ss）；失败时回退系统字体。"""
    if "fb" not in _TEXT_CACHE:
        data = None
        try:
            data = fitz.Font(FONT_CJK).buffer
        except Exception:
            data = None
        if not data:
            for p in (r"C:\Windows\Fonts\simhei.ttf",
                      r"C:\Windows\Fonts\msyh.ttc",
                      r"C:\Windows\Fonts\simsun.ttc"):
                if os.path.exists(p):
                    try:
                        with open(p, "rb") as f:
                            data = f.read()
                    except OSError:
                        data = None
                    if data:
                        break
        _TEXT_CACHE["fb"] = data
    return _TEXT_CACHE["fb"]


def _pil_font(size_px):
    from PIL import ImageFont
    key = ("f", int(size_px))
    if key not in _TEXT_CACHE:
        data = _cjk_font_bytes()
        if data:
            _TEXT_CACHE[key] = ImageFont.truetype(io.BytesIO(data), int(size_px))
        else:
            _TEXT_CACHE[key] = ImageFont.load_default()
    return _TEXT_CACHE[key]


def text_image(text, fs, color, angle=0):
    """把文字渲染为透明背景 PIL 图像（1 像素 = 1/TEXT_SS 点）。
    color 为 (r,g,b) 浮点 0~1；angle 为顺时针角度。返回 PIL.Image。"""
    from PIL import Image, ImageDraw
    font = _pil_font(fs * TEXT_SS)
    l, t, r, b = font.getbbox(text)
    w, h = max(1, int(r - l)), max(1, int(b - t))
    img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    fill = (int(color[0] * 255), int(color[1] * 255), int(color[2] * 255), 255)
    ImageDraw.Draw(img).text((-l, -t), text, font=font, fill=fill)
    if angle:
        img = img.rotate(-angle, expand=True, resample=Image.BICUBIC)
    return img


def text_png(text, fs, color, angle=0):
    """同 text_image，但返回 (PNG 字节, 宽px, 高px)。"""
    img = text_image(text, fs, color, angle)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue(), img.width, img.height

CONFIG_DIR = os.path.join(os.environ.get("APPDATA", os.path.expanduser("~")), "3Q PDF开源免费工具")
CONFIG_PATH = os.path.join(CONFIG_DIR, "config.json")


def open_pdf_doc(path, password=None):
    """打开 PDF；带密码时自动认证。失败抛出 ValueError(中文说明)。"""
    doc = fitz.open(path)
    if doc.needs_pass:
        if not password or not doc.authenticate(password):
            doc.close()
            raise ValueError("文件已加密且密码不正确，无法处理。")
    return doc


# ============================================================
# 核心处理层（不依赖界面，可独立测试）
# ============================================================

def parse_page_expr(expr, total):
    """解析页码表达式，如 '3,5,12-18,40'（支持中文逗号）。
    返回 0 基页码列表（按书写顺序，允许重复）。非法输入抛出 ValueError(中文说明)。"""
    if not expr or not expr.strip():
        raise ValueError("页码不能为空")
    pages = []
    for seg in expr.replace("，", ",").split(","):
        seg = seg.strip()
        if not seg:
            continue
        if "-" in seg:
            parts = seg.split("-")
            if len(parts) != 2 or not parts[0].strip().isdigit() or not parts[1].strip().isdigit():
                raise ValueError(f"无法识别的页码段：“{seg}”")
            a, b = int(parts[0]), int(parts[1])
            if a > b:
                raise ValueError(f"页码范围“{seg}”起页大于止页")
            if a < 1 or b > total:
                raise ValueError(f"页码范围“{seg}”超出总页数（共 {total} 页）")
            pages.extend(range(a - 1, b))
        else:
            if not seg.isdigit():
                raise ValueError(f"无法识别的页码：“{seg}”")
            p = int(seg)
            if p < 1 or p > total:
                raise ValueError(f"页码 {p} 超出总页数（共 {total} 页）")
            pages.append(p - 1)
    if not pages:
        raise ValueError("页码不能为空")
    return pages


def parse_ranges_expr(expr, total):
    """解析拆分用的页码范围组，如 '1-10,11-30,31-60'。
    返回 [(start0, end0), ...]（0 基，含端点）。范围不得重叠。"""
    pages_groups = []
    seen = set()
    for seg in expr.replace("，", ",").split(","):
        seg = seg.strip()
        if not seg:
            continue
        if "-" in seg:
            parts = seg.split("-")
            if len(parts) != 2 or not parts[0].strip().isdigit() or not parts[1].strip().isdigit():
                raise ValueError(f"无法识别的页码段：“{seg}”")
            a, b = int(parts[0]), int(parts[1])
            if a > b:
                raise ValueError(f"页码范围“{seg}”起页大于止页")
            if a < 1 or b > total:
                raise ValueError(f"页码范围“{seg}”超出总页数（共 {total} 页）")
        else:
            if not seg.isdigit():
                raise ValueError(f"无法识别的页码：“{seg}”")
            a = b = int(seg)
            if a < 1 or a > total:
                raise ValueError(f"页码 {a} 超出总页数（共 {total} 页）")
        grp = list(range(a - 1, b))
        dup = seen.intersection(grp)
        if dup:
            raise ValueError(f"第 {min(dup) + 1} 页在多个范围中重复出现")
        seen.update(grp)
        pages_groups.append((a - 1, b - 1))
    if not pages_groups:
        raise ValueError("页码范围不能为空")
    return pages_groups


def unique_out_path(path):
    """若文件已存在，自动追加序号 (2)(3)…，保证不覆盖（SR-01）。"""
    if not os.path.exists(path):
        return path
    base, ext = os.path.splitext(path)
    i = 2
    while os.path.exists(f"{base} ({i}){ext}"):
        i += 1
    return f"{base} ({i}){ext}"


def check_pdf_readable(path):
    """检测 PDF 是否可读（SR-03）。返回 (是否可读, 提示信息, 页数)。"""
    try:
        doc = fitz.open(path)
        if doc.needs_pass:
            doc.close()
            return False, "文件已加密", 0
        n = doc.page_count
        doc.close()
        return True, "", n
    except Exception:
        return False, "文件损坏或不是有效的 PDF", 0

def do_merge(files, out_path, progress=None):
    """FR-01 合并：按列表顺序拼接。"""
    out = fitz.open()
    total = 0
    for i, f in enumerate(files):
        src = fitz.open(f)
        try:
            out.insert_pdf(src)
            total += src.page_count
        finally:
            src.close()
        if progress:
            progress(i + 1, len(files))
    out.save(out_path, garbage=3, deflate=True)
    out.close()
    return total


def do_split(src_path, ranges, out_dir, base_name, pw=None):
    """FR-02 拆分：ranges 为 [(start0, end0), ...]，返回输出文件路径列表。"""
    os.makedirs(out_dir, exist_ok=True)
    src = open_pdf_doc(src_path, pw)
    outs = []
    for s0, e0 in ranges:
        out = fitz.open()
        out.insert_pdf(src, from_page=s0, to_page=e0)
        p = unique_out_path(os.path.join(out_dir, f"{base_name}_第{s0 + 1}-{e0 + 1}页.pdf"))
        out.save(p, garbage=3, deflate=True)
        out.close()
        outs.append(p)
    src.close()
    return outs


def do_extract(src_path, pages, out_path, pw=None):
    """FR-03 提取：按给定顺序取页合成一个新文件。"""
    src = open_pdf_doc(src_path, pw)
    out = fitz.open()
    for p in pages:
        out.insert_pdf(src, from_page=p, to_page=p)
    out.save(out_path, garbage=3, deflate=True)
    out.close()
    src.close()


def text_width(text, fontsize):
    """内置中文字体的真实排字宽度（与 insert_text 实际落墨一致，单位：点）。
    实测（2026-10-17）：fitz.Font("china-ss").text_length 对 "第 1 页 / 共 4 页"@10pt
    只给 70.47，而实际落墨 130.0（每字全角）；ASCII "Page 1 of 4" 给 49.57，
    实际 110.0。用它算居中会偏约 1.8 倍宽度（页码右移 ~30pt），长文字也不会触发缩字号。
    fitz.get_text_length 与内置字体同源，量得 130.0 / 110.0，与落墨逐点吻合。"""
    try:
        return fitz.get_text_length(text, fontname=FONT_CJK, fontsize=fontsize)
    except Exception:
        return fitz.Font(FONT_CJK).text_length(text, fontsize)


def _fit_fontsize(font, text, fontsize, max_width):
    """文字超宽时自动缩小字号（最小 6）。font 参数保留兼容，宽度以 text_width 为准。"""
    while fontsize > 6 and text_width(text, fontsize) > max_width:
        fontsize -= 1
    return fontsize


def do_insert_text(src_path, out_path, page_indices, text, fontsize, color,
                   pos_preset, custom_xy_mm=None, angle=0, pw=None, vector=False):
    """FR-04 插入文字：叠加到页面最上层。angle 为顺时针角度（度，与预览一致）。
    实测（2026-10）：morph Matrix(+45) 输出 dir=(cos45,-sin45)，显示坐标下为逆时针；
    故传给 morph 前取负，使正角度=顺时针。返回输出路径。"""
    font = fitz.Font(FONT_CJK)
    doc = open_pdf_doc(src_path, pw)
    for pno in page_indices:
        page = doc[pno]
        r = page.rect
        fs = _fit_fontsize(font, text, fontsize, r.width * 0.9)
        w = text_width(text, fs)
        mat = fitz.Matrix(-angle) if angle else None

        def put(x, y, anchor=None):
            if mat is not None:
                page.insert_text(fitz.Point(x, y), text, fontname=FONT_CJK,
                                 fontsize=fs, color=color,
                                 morph=(fitz.Point(anchor), mat))
            else:
                page.insert_text(fitz.Point(x, y), text, fontname=FONT_CJK,
                                 fontsize=fs, color=color)

        if pos_preset == "页面居中":
            cx, cy = r.width / 2, r.height / 2
            put(cx - w / 2, cy, anchor=(cx, cy))
        elif pos_preset == "页眉":
            put((r.width - w) / 2, 20 * MM)
        elif pos_preset == "页脚":
            put((r.width - w) / 2, r.height - 12 * MM)
        elif pos_preset == "四角":
            m = 12 * MM
            for x, y in [(m, m + fs), (r.width - m - w, m + fs),
                         (m, r.height - m), (r.width - m - w, r.height - m)]:
                put(x, y)
        else:  # 自定义坐标（毫米，原点页面左上角）：以图片贴入，位置/大小所见即所得
            x_mm, y_mm = custom_xy_mm
            if vector:
                # 可选文本：矢量插入（文字可搜索/复制）。用与贴纸完全相同的字体与
                # 字形包围盒换算落点，保证矢量文字与预览位置一致；有角度时绕字形中心旋转。
                pfont = _pil_font(fs * TEXT_SS)
                l, t, r2, b = pfont.getbbox(text)
                try:
                    asc = pfont.getmetrics()[0]  # 基线距布局顶的距离（像素）
                except Exception:
                    asc = b
                x, y = x_mm * MM, y_mm * MM
                gw, gh = (r2 - l) / TEXT_SS, (b - t) / TEXT_SS
                base_img = text_image(text, fs, color, angle)
                cw, ch = base_img.width / TEXT_SS, base_img.height / TEXT_SS
                cx, cy = x + cw / 2, y + ch / 2      # 字形中心（=旋转中心）
                xg, yg = cx - gw / 2, cy - gh / 2    # 未旋转字形左上角
                # 基线：贴纸字形顶在 yg，基线在字形顶下方 (asc - t) 处
                bx, by = xg - l / TEXT_SS, yg + (asc - t) / TEXT_SS
                # 必须嵌入与贴纸完全相同的字体（Droid Sans Fallback），否则
                # 「china-ss」会被替换为内置宋体，字宽与预览不一致。
                fname = "pdfsticker"
                try:
                    page.insert_font(fontname=fname, fontbuffer=_cjk_font_bytes())
                except Exception:
                    fname = FONT_CJK
                if angle:
                    page.insert_text(fitz.Point(bx, by), text, fontname=fname,
                                     fontsize=fs, color=color,
                                     morph=(fitz.Point(cx, cy), fitz.Matrix(-angle)))
                else:
                    page.insert_text(fitz.Point(bx, by), text, fontname=fname,
                                     fontsize=fs, color=color)
            else:
                stream, wpx, hpx = text_png(text, fs, color, angle)
                wpt, hpt = wpx / TEXT_SS, hpx / TEXT_SS
                rect = fitz.Rect(x_mm * MM, y_mm * MM,
                                 x_mm * MM + wpt, y_mm * MM + hpt)
                page.insert_image(rect, stream=stream)
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return out_path


def do_insert_image(src_path, out_path, img_path, page_indices,
                    pos_preset, custom_xy_mm=None, width_pct=20, rotate=0, pw=None):
    """FR-05 插入图片：width_pct 为相对图片「物理尺寸」的缩放百分比（100=按图片
    自身 DPI 还原的真实大小，无 DPI 时按 96dpi），等比缩放。rotate 仅支持 90 度
    倍数（顺时针）。返回输出路径。"""
    pix = fitz.Pixmap(img_path)
    iw, ih = pix.width, pix.height
    pix = None
    size = image_physical_size(img_path) or (iw * 72 / 96.0, ih * 72 / 96.0)
    base_w, base_h = size
    doc = open_pdf_doc(src_path, pw)
    for pno in page_indices:
        page = doc[pno]
        dw, dh = page_geometry(page)[:2]
        w = base_w * width_pct / 100.0
        h = base_h * width_pct / 100.0
        if rotate in (90, 270):
            w, h = h, w  # 旋转后占位宽高互换
        if pos_preset == "右下角":
            m = 10 * MM
            dx0, dy0 = dw - m - w, dh - m - h
        elif pos_preset == "页面居中":
            dx0, dy0 = (dw - w) / 2, (dh - h) / 2
        else:  # 自定义坐标（毫米，图片占位框左上角，显示坐标系）
            x_mm, y_mm = custom_xy_mm
            dx0, dy0 = x_mm * MM, y_mm * MM
        # insert_image 用未旋转坐标系：把显示坐标的左上角反算回去
        _dw, _dh, rot, mbw, mbh = page_geometry(page)
        ux, uy = _rot_unmap(rot, mbw, mbh, dx0, dy0)
        if rot in (90, 270):
            rect = fitz.Rect(ux - h, uy, ux, uy + w)  # 占位框宽高互换
        else:
            rect = fitz.Rect(ux, uy, ux + w, uy + h)
        page.insert_image(rect, filename=img_path, rotate=rotate)
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return out_path


# ---------------- 新增核心功能（V1.3） ----------------

def do_images_to_pdf(image_paths, out_path, fit_a4=False, progress=None):
    """图片转 PDF：每张图片一页。fit_a4=False 时页面尺寸按图片自身 DPI 还原
    （300dpi 扫描件 = 真实 A4；无 DPI 信息则按 96dpi 屏幕尺寸）；
    True 则放入 A4 纸并留边距居中（适合打印）。"""
    out = fitz.open()
    a4w, a4h = 595.0, 842.0
    for i, p in enumerate(image_paths):
        pix = fitz.Pixmap(p)
        iw, ih = pix.width, pix.height
        pix = None
        if fit_a4:
            page = out.new_page(width=a4w, height=a4h)
            m = 10 * MM
            s = min((a4w - 2 * m) / iw, (a4h - 2 * m) / ih)
            w, h = iw * s, ih * s
            rect = fitz.Rect((a4w - w) / 2, (a4h - h) / 2, (a4w + w) / 2, (a4h + h) / 2)
        else:
            size = image_physical_size(p)
            if size:
                pw_pt, ph_pt = size
            else:  # 取不到 DPI 时退回 96dpi
                pw_pt, ph_pt = iw * 72 / 96.0, ih * 72 / 96.0
            page = out.new_page(width=pw_pt, height=ph_pt)
            rect = page.rect
        page.insert_image(rect, filename=p)
        if progress:
            progress(i + 1, len(image_paths))
    out.save(out_path, garbage=3, deflate=True)
    out.close()


def do_pdf_to_images(src_path, out_dir, fmt="png", dpi=150, pw=None, progress=None):
    """PDF 转图片：每页导出为 PNG/JPG。返回输出路径列表。"""
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.splitext(os.path.basename(src_path))[0]
    doc = open_pdf_doc(src_path, pw)
    outs = []
    for i in range(doc.page_count):
        pix = doc[i].get_pixmap(dpi=dpi, alpha=False)
        p = unique_out_path(os.path.join(out_dir, f"{base}_第{i + 1}页.{fmt}"))
        pix.save(p)
        outs.append(p)
        if progress:
            progress(i + 1, doc.page_count)
    doc.close()
    return outs


def do_delete_pages(src_path, pages0, out_path, pw=None):
    """删除指定页（0 基列表），其余页保持原顺序，另存新文件。"""
    doc = open_pdf_doc(src_path, pw)
    keep = [i for i in range(doc.page_count) if i not in set(pages0)]
    if not keep:
        doc.close()
        raise ValueError("不能删除全部页面，至少要保留 1 页。")
    doc.select(keep)
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return len(keep)


def do_rotate_pages(src_path, pages0, angle, out_path, pw=None):
    """永久旋转指定页（angle 为顺时针 90/180/270），另存新文件。"""
    doc = open_pdf_doc(src_path, pw)
    for p in pages0:
        doc[p].set_rotation((doc[p].rotation + angle) % 360)
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()


def do_add_page_numbers(src_path, out_path, position="底部居中", start=1,
                        fmt="第 X 页 / 共 N 页", fontsize=10, pw=None):
    """添加页码。fmt 中 X=当前页号，N=总页数。
    定位在「显示坐标系」中计算，旋转页（/Rotate 90/180/270）也落在视觉正确位置。"""
    doc = open_pdf_doc(src_path, pw)
    font = fitz.Font(FONT_CJK)
    n = doc.page_count
    for i, page in enumerate(doc):
        s = fmt.replace("X", str(start + i)).replace("N", str(n))
        dw, dh = page_geometry(page)[:2]
        fs = _fit_fontsize(font, s, fontsize, dw * 0.9)
        w = text_width(s, fs)
        if position == "底部居中":
            x, y = (dw - w) / 2, dh - 10 * MM
        elif position == "底部右侧":
            x, y = dw - w - 12 * MM, dh - 10 * MM
        elif position == "底部左侧":
            x, y = 12 * MM, dh - 10 * MM
        elif position == "顶部右侧":
            x, y = dw - w - 12 * MM, 12 * MM + fs
        else:  # 顶部左侧
            x, y = 12 * MM, 12 * MM + fs
        insert_text_display(page, x, y, s, FONT_CJK, fs, (0, 0, 0))
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()


def _image_display_rects(doc):
    """返回 {图片 xref: 最大显示矩形(pt)}。
    实测（2026-10-17，PyMuPDF 1.25.2）：page.get_image_info() 的 'xref' 恒为 None，
    必须走 page.get_images() 拿 xref，再用 page.get_image_rects(xref) 拿显示框。"""
    rects = {}
    for pno in range(doc.page_count):
        try:
            items = doc[pno].get_images(full=True)
        except Exception:
            continue
        for item in items:
            x = item[0]
            if not x:
                continue
            try:
                for r in doc[pno].get_image_rects(x):
                    a = abs(r.width * r.height)
                    cur = rects.get(x)
                    if cur is None or a > abs(cur.width * cur.height):
                        rects[x] = fitz.Rect(r)
            except Exception:
                continue
    return rects


def _recompress_images(doc, quality=50, max_dpi=0):
    """逐张图片重压：显示分辨率高于 max_dpi 就降采样，再按 quality 重编为 JPEG。
    max_dpi=0 表示不动分辨率、只重编码。返回 (改动张数, 跳过张数)。
    实测（2026-10-17）：PyMuPDF 1.25.2 根本没有 Document.rewrite_images（AttributeError），
    老代码用 hasattr 挡着、异常被 pass 掉，等于一直没压。
    自建方案实测 5.33MB（3 张 1600x1200 噪声图）→ q50/dpi150 = 0.63MB，q30/dpi96 = 0.11MB。
    跳过：被别人当 SMask 用的蒙版（降采样会和底图错位）、无 SMask 的透明图（转 JPEG 变黑块）、
    1 位黑白蒙版、以及重编后反而变大的图。"""
    from PIL import Image as PILImage
    smasks = set()
    for x in range(1, doc.xref_length()):
        try:
            k = doc.xref_get_key(x, "SMask")
        except Exception:
            continue
        if k and k[0] == "xref":
            try:
                smasks.add(int(str(k[1]).split()[0]))
            except Exception:
                pass
    rects = _image_display_rects(doc)
    done = skipped = 0
    for x in range(1, doc.xref_length()):
        if x in smasks:
            continue
        try:
            if doc.xref_get_key(x, "Subtype") != ("name", "/Image"):
                continue
            if doc.xref_get_key(x, "ImageMask")[0] == "bool":     # 纯蒙版
                continue
            bpc = doc.xref_get_key(x, "BitsPerComponent")
            if bpc[0] == "int" and int(bpc[1]) <= 1:              # 1 位蒙版
                continue
        except Exception:
            continue
        try:
            raw = doc.extract_image(x)
            with PILImage.open(io.BytesIO(raw["image"])) as src:
                im = src.copy()
            if im.width < 8 or im.height < 8:
                skipped += 1
                continue
            r = rects.get(x)
            eff = 0.0
            if r and r.width > 1 and r.height > 1:
                eff = max(im.width / (r.width / 72.0), im.height / (r.height / 72.0))
            if max_dpi and eff > max_dpi:
                s = max_dpi / eff
                im = im.resize((max(1, int(im.width * s)), max(1, int(im.height * s))),
                               PILImage.LANCZOS)
            if im.mode in ("RGBA", "LA", "P"):
                # 透明图：alpha 在 SMask 里才能压，否则转 JPEG 会把透明底压成黑块
                has_sm = False
                try:
                    has_sm = doc.xref_get_key(x, "SMask")[0] == "xref"
                except Exception:
                    pass
                if not has_sm:
                    skipped += 1
                    continue
            im = im.convert("RGB")
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=int(max(5, min(95, quality))), optimize=True)
            data = buf.getvalue()
            if len(data) >= len(raw["image"]) * 0.98:
                skipped += 1
                continue
            doc.update_stream(x, data)
            doc.xref_set_key(x, "Filter", "/DCTDecode")
            doc.xref_set_key(x, "ColorSpace", "/DeviceRGB")
            doc.xref_set_key(x, "BitsPerComponent", "8")
            doc.xref_set_key(x, "DecodeParms", "null")
            done += 1
        except Exception:
            skipped += 1
    return done, skipped


def do_compress(src_path, out_path, quality=50, max_dpi=0, pw=None, progress=None):
    """压缩体积：逐张图片按 quality 重编码为 JPEG，可选把显示分辨率压到 max_dpi 以下。
    max_dpi=0 表示只重编码、不动分辨率。返回 (原大小, 新大小) 字节。"""
    old = os.path.getsize(src_path)
    doc = open_pdf_doc(src_path, pw)
    _recompress_images(doc, quality, max_dpi)
    doc.save(out_path, garbage=4, deflate=True, clean=True)
    doc.close()
    return old, os.path.getsize(out_path)


# ---------------- 压缩到指定大小 ----------------

def do_compress_to_target(src_path, out_path, target_bytes, pw=None, progress=None):
    """压到目标体积以内：先二分画质，不够再二分分辨率，最后在低分辨率下再二分画质。
    返回 (原大小, 新大小, 是否达成目标)。
    文字型 PDF（没有图片可压）压不下去时会如实返回 False，取到的仍是最小的那一份。"""
    old = os.path.getsize(src_path)
    tmpd = tempfile.mkdtemp(prefix="cmpt_")
    tried = []                       # [(大小, 临时路径, max_dpi, quality)]
    n = [0]

    def best_so_far():
        return min([t[0] for t in tried] + [old])

    def attempt(dpi, q):
        n[0] += 1
        if progress:
            progress(n[0], 18)
        p = os.path.join(tmpd, "t_%d_%d.pdf" % (int(dpi), int(q)))
        doc = open_pdf_doc(src_path, pw)
        _recompress_images(doc, q, dpi)
        doc.save(p, garbage=4, deflate=True, clean=True)
        doc.close()
        s = os.path.getsize(p)
        tried.append((s, p, dpi, q))
        return s

    def bisect(lo, hi, make, iters):
        """体积随参数单调递增，找「仍能装进目标的最大参数值」。"""
        for _ in range(iters):
            if lo > hi:
                break
            mid = (lo + hi) // 2
            if make(mid) <= target_bytes:
                lo = mid + 1
            else:
                hi = mid - 1

    try:
        # 阶段一：不动分辨率，只二分画质（先保住原始清晰度）
        bisect(8, 85, lambda q: attempt(0, q), 6)
        # 阶段二：画质到底仍不达标 —— 二分分辨率上限
        if best_so_far() > target_bytes:
            bisect(36, 150, lambda d: attempt(d, 45), 6)
        # 阶段三：最低分辨率下再二分画质
        if best_so_far() > target_bytes:
            bisect(5, 45, lambda q: attempt(36, q), 5)
        fits = [t for t in tried if t[0] <= target_bytes]
        if not fits and old <= target_bytes:
            # 本来就比目标小：给一份清理过的副本
            p = os.path.join(tmpd, "orig.pdf")
            doc = open_pdf_doc(src_path, pw)
            doc.save(p, garbage=4, deflate=True, clean=True)
            doc.close()
            tried.append((os.path.getsize(p), p, 0, 85))
            fits = [t for t in tried if t[0] <= target_bytes]
        # 达标时取「仍能装进目标的最大体积」= 最高画质；不达标则取最小的一份
        best = max(fits, key=lambda t: t[0]) if fits else min(tried, key=lambda t: t[0])
        with open(best[1], "rb") as f:
            data = f.read()
        with open(out_path, "wb") as f:
            f.write(data)
        return old, len(data), bool(fits)
    finally:
        try:
            for f in os.listdir(tmpd):
                os.remove(os.path.join(tmpd, f))
            os.rmdir(tmpd)
        except OSError:
            pass


# ---------------- Office → PDF（Windows + 本机 Office） ----------------

OFFICE_EXTS = (".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx", ".csv")


def office_kind(path):
    """按扩展名判断该用哪个 Office 组件；不支持返回 None。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".doc", ".docx"):
        return "word"
    if ext in (".xls", ".xlsx", ".csv"):
        return "excel"
    if ext in (".ppt", ".pptx"):
        return "ppt"
    return None


def office_to_pdf(src_path, out_path):
    """Word / Excel / PowerPoint 转 PDF。依赖 Windows + 本机已装对应 Office 组件（COM）。
    缺 pywin32 或没装 Office 时抛带处置建议的异常，不静默失败。"""
    kind = office_kind(src_path)
    if kind is None:
        raise ValueError("不支持的文件类型：%s（支持 doc/docx/xls/xlsx/ppt/pptx/csv）"
                         % os.path.basename(src_path))
    if os.name != "nt":
        raise RuntimeError("Office 转 PDF 走 Windows COM，当前系统不支持。")
    try:
        import win32com.client as win32
    except Exception:
        raise RuntimeError("缺少依赖 pywin32：请在命令行执行 pip install pywin32 后重试。")
    name = {"word": "Word", "excel": "Excel", "ppt": "PowerPoint"}[kind]
    src = os.path.abspath(src_path)
    out = os.path.abspath(out_path)
    d = os.path.dirname(out)
    if d:
        os.makedirs(d, exist_ok=True)
    app = None
    try:
        if kind == "word":
            app = win32.Dispatch("Word.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            doc = app.Documents.Open(src, ReadOnly=True)
            try:
                doc.SaveAs(out, FileFormat=17)          # wdFormatPDF
            finally:
                doc.Close(False)
        elif kind == "excel":
            app = win32.Dispatch("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            wb = app.Workbooks.Open(src, ReadOnly=True)
            try:
                wb.ExportAsFixedFormat(0, out)          # xlTypePDF
            finally:
                wb.Close(False)
        else:
            app = win32.Dispatch("PowerPoint.Application")
            pres = app.Presentations.Open(src, True, False, False)
            try:
                pres.SaveAs(out, 32)                    # ppSaveAsPDF
            finally:
                pres.Close()
    except Exception as e:
        raise RuntimeError("用 %s 转换失败：%s\n请确认本机已安装 %s（WPS 需开启 COM 兼容）。"
                           % (name, e, name))
    finally:
        if app is not None:
            try:
                app.Quit()
            except Exception:
                pass
    if not os.path.exists(out):
        raise RuntimeError("%s 未生成 PDF：%s" % (name, os.path.basename(src_path)))
    return out


def office_batch_to_pdf(files, out_dir, progress=None):
    """批量 Office → PDF。返回 (成功输出列表, [(文件名, 失败原因), ...])。"""
    os.makedirs(out_dir, exist_ok=True)
    outs, errs = [], []
    for i, f in enumerate(files):
        base = os.path.splitext(os.path.basename(f))[0]
        out = unique_out_path(os.path.join(out_dir, base + ".pdf"))
        try:
            office_to_pdf(f, out)
            outs.append(out)
        except Exception as e:
            errs.append((os.path.basename(f), str(e)))
        if progress:
            progress(i + 1, len(files))
    return outs, errs


# ---------------- 页面裁剪 / 去黑边 ----------------

def page_content_bbox(page, dpi=100, white=240):
    """渲染页面求「非白内容」包围盒，返回未旋转坐标系的 fitz.Rect；整页空白返回 None。
    扫描件黑边是深色像素，阈值 240 能切得干净又不伤正文（实测见 2026-10-17 回归）。"""
    from PIL import Image as PILImage
    pix = page.get_pixmap(dpi=dpi, alpha=False)
    mode = {1: "L", 3: "RGB", 4: "RGBA"}.get(pix.n, "RGB")
    try:
        im = PILImage.frombytes(mode, (pix.width, pix.height), pix.samples)
    except Exception:
        im = PILImage.frombytes("RGB", (pix.width, pix.height), pix.samples)
    mask = im.convert("L").point(lambda v: 255 if v < white else 0)
    bb = mask.getbbox()
    if not bb:
        return None
    s = 72.0 / dpi
    _dw, _dh, rot, mbw, mbh = page_geometry(page)
    x0, y0, x1, y1 = [v * s for v in bb]
    pts = [_rot_unmap(rot, mbw, mbh, x0, y0), _rot_unmap(rot, mbw, mbh, x1, y0),
           _rot_unmap(rot, mbw, mbh, x0, y1), _rot_unmap(rot, mbw, mbh, x1, y1)]
    xs = [q[0] for q in pts]
    ys = [q[1] for q in pts]
    return fitz.Rect(min(xs), min(ys), max(xs), max(ys))


def do_crop_pages(src_path, out_path, pages=None, mode="自动去黑边", keep_mm=3.0,
                  margins_mm=(0, 0, 0, 0), pw=None, progress=None):
    """裁剪页面白边/黑边。mode：自动去黑边 / 四边留白 / 恢复原始尺寸。
    margins_mm = (左, 上, 右, 下) 内缩量。返回 (处理页数, 实际改动页数)。"""
    doc = open_pdf_doc(src_path, pw)
    page_list = list(range(doc.page_count)) if pages is None else list(pages)
    keep = float(keep_mm) * MM
    changed = 0
    for k, pno in enumerate(page_list):
        page = doc[pno]
        mb = page.mediabox
        box = None
        if mode == "恢复原始尺寸":
            page.set_cropbox(fitz.Rect(mb))
            changed += 1
        elif mode == "四边留白":
            l, t, r, b = [float(v) * MM for v in (margins_mm or (0, 0, 0, 0))]
            box = fitz.Rect(mb.x0 + l, mb.y0 + t, mb.x1 - r, mb.y1 - b)
        else:
            bb = page_content_bbox(page)
            if bb is not None:
                box = fitz.Rect(max(mb.x0, bb.x0 - keep), max(mb.y0, bb.y0 - keep),
                                min(mb.x1, bb.x1 + keep), min(mb.y1, bb.y1 + keep))
        if box is not None:
            # 内容几乎铺满整页就没得裁；裁完不足 1cm 视为误判
            if (box.width > 28 and box.height > 28
                    and box.get_area() < mb.get_area() * 0.995
                    and (box.x0 > mb.x0 or box.y0 > mb.y0
                         or box.x1 < mb.x1 or box.y1 < mb.y1)):
                page.set_cropbox(box)
                changed += 1
        if progress:
            progress(k + 1, len(page_list))
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return len(page_list), changed


# ---------------- 图片预处理（手机照片方向 / 旋转） ----------------

def image_exif_rotation(path):
    """手机照片 EXIF 方向对应的「顺时针」角度（0/90/180/270）。
    实测：MuPDF 贴图不应用 EXIF 方向，竖拍照片会横躺，必须自己转。"""
    from PIL import Image as PILImage
    try:
        with PILImage.open(path) as im:
            o = im.getexif().get(274, 1)        # 0x0112 Orientation
    except Exception:
        return 0
    return {3: 180, 6: 90, 8: 270}.get(int(o or 1), 0)


def rotate_image_file(path, angle):
    """把图片顺时针旋转 angle 度后另存为临时 PNG（保留 DPI）。返回临时文件路径。
    调用方负责在不再需要时删除。"""
    from PIL import Image as PILImage
    a = int(angle or 0) % 360
    with PILImage.open(path) as src:
        im = src.copy()
        dpi = src.info.get("dpi")
    if a == 90:
        im = im.transpose(PILImage.ROTATE_270)
    elif a == 180:
        im = im.transpose(PILImage.ROTATE_180)
    elif a == 270:
        im = im.transpose(PILImage.ROTATE_90)
    im = im.convert("RGBA") if im.mode in ("RGBA", "LA", "P") else im.convert("RGB")
    fd, tmp = tempfile.mkstemp(suffix=".png", prefix="rot_")
    os.close(fd)
    try:
        im.save(tmp, "PNG", dpi=tuple(dpi) if dpi else (96, 96))
    except Exception:
        im.save(tmp, "PNG")
    return tmp


def do_encrypt(src_path, out_path, user_pw, pw=None):
    """加密：设置打开密码（AES-256），保留打印/复制/无障碍权限。"""
    doc = open_pdf_doc(src_path, pw)
    doc.save(out_path, garbage=3, deflate=True,
             encryption=fitz.PDF_ENCRYPT_AES_256,
             user_pw=user_pw, owner_pw=user_pw,
             permissions=fitz.PDF_PERM_ACCESSIBILITY | fitz.PDF_PERM_PRINT
             | fitz.PDF_PERM_COPY)
    doc.close()


def do_decrypt(src_path, out_path, pw=None):
    """解密：去除打开密码与权限限制，另存为无加密副本。"""
    doc = open_pdf_doc(src_path, pw)
    doc.save(out_path, garbage=3, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)
    doc.close()
    return out_path


# ---------------- 水印 ----------------

def _apply_alpha(img, opacity):
    """按不透明度比例缩放图像的整体 alpha 通道（opacity 0~1）。"""
    if opacity >= 1.0:
        return img
    img = img.copy()
    a = img.getchannel("A").point(lambda v: int(v * opacity))
    img.putalpha(a)
    return img


def _watermark_tile(pw, ph, w, h, layout, gap_x=25.0, gap_y=40.0):
    """计算水印放置位置（左上角坐标，单位点）。layout 为「居中一个」或「平铺整页」。"""
    if layout == "居中一个":
        return [((pw - w) / 2, (ph - h) / 2)]
    step_x, step_y = w + gap_x, h + gap_y
    xs, x = [], -w
    while x < pw:
        xs.append(x)
        x += step_x
    ys, y = [], -h
    while y < ph:
        ys.append(y)
        y += step_y
    return [(x, y) for y in ys for x in xs]


def do_watermark_text(src_path, out_path, pages, text, fontsize, color, angle,
                      opacity, layout, pw=None):
    """文字水印：用带透明度与角度的透明 PNG 贴入页面（支持平铺/居中）。"""
    img = _apply_alpha(text_image(text, fontsize, color, angle), opacity)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    stream = buf.getvalue()
    w, h = img.width / TEXT_SS, img.height / TEXT_SS
    doc = open_pdf_doc(src_path, pw)
    if pages is None:
        pages = range(doc.page_count)
    for pno in pages:
        page = doc[pno]
        mb = page.mediabox
        for x, y in _watermark_tile(mb.width, mb.height, w, h, layout):
            page.insert_image(fitz.Rect(x, y, x + w, y + h), stream=stream)
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return out_path


def do_watermark_image(src_path, out_path, pages, img_path, scale_pct, opacity,
                       layout, angle=0, pw=None):
    """图片水印：等比缩放后贴入页面（支持透明度、角度、平铺/居中）。"""
    from PIL import Image as PILImage
    pil = PILImage.open(img_path).convert("RGBA")
    pil = _apply_alpha(pil, opacity)
    if angle:
        pil = pil.rotate(-angle, expand=True, resample=PILImage.BICUBIC)
    buf = io.BytesIO()
    pil.save(buf, "PNG")
    stream = buf.getvalue()
    s = scale_pct / 100.0
    w, h = pil.width * s, pil.height * s
    doc = open_pdf_doc(src_path, pw)
    if pages is None:
        pages = range(doc.page_count)
    for pno in pages:
        page = doc[pno]
        mb = page.mediabox
        for x, y in _watermark_tile(mb.width, mb.height, w, h, layout):
            page.insert_image(fitz.Rect(x, y, x + w, y + h), stream=stream)
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return out_path


def do_reorder(src_path, out_path, plan, pw=None):
    """页面管理：按 plan 重建文档。plan 元素为 ('page', 0基页号) 或
    ('blank', 宽pt, 高pt)，可移动/复制/删除/插入空白页，结果另存新文件。"""
    src = open_pdf_doc(src_path, pw)
    out = fitz.open()
    for item in plan:
        if item[0] == "page":
            out.insert_pdf(src, from_page=item[1], to_page=item[1])
        else:
            out.new_page(width=item[1], height=item[2])
    out.save(out_path, garbage=3, deflate=True)
    out.close()
    src.close()
    return len(plan)


# ============================================================
# V1.5 核心层：格式转换 / 书签 / 表单 / 批量
# ============================================================

def do_pdf_to_docx(src_path, out_path, pages=None, pw=None, progress=None):
    """PDF 转 Word（.docx）：抽取文字块与图片，按页生成段落。
    适配文字型/图文型 PDF；复杂版式（多栏、表格、公式）保真度有限。"""
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.oxml.ns import qn

    doc = open_pdf_doc(src_path, pw)
    page_list = list(range(doc.page_count)) if pages is None else list(pages)
    if not page_list:
        doc.close()
        raise ValueError("未选择任何页面。")
    out = Document()
    try:
        rpr = out.styles["Normal"].element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    except Exception:
        pass

    def add_run(p, text, size, bold):
        run = p.add_run(text)
        run.font.name = "宋体"
        try:
            run.font.size = Pt(min(72, max(5, float(size))))
        except Exception:
            pass
        run.font.bold = bool(bold)
        try:
            run._element.get_or_add_rPr().get_or_add_rFonts().set(
                qn("w:eastAsia"), "宋体")
        except Exception:
            pass

    for i, pno in enumerate(page_list):
        page = doc[pno]
        if i > 0:
            out.add_page_break()
        data = page.get_text("dict")
        for block in data.get("blocks", []):
            btype = block.get("type")
            if btype == 0:
                for line in block.get("lines", []):
                    spans = line.get("spans", [])
                    if not spans:
                        continue
                    p = out.add_paragraph()
                    try:
                        p.paragraph_format.space_after = Pt(2)
                    except Exception:
                        pass
                    for s in spans:
                        add_run(p, s.get("text", ""), s.get("size", 10.5),
                                "Bold" in (s.get("font") or ""))
            elif btype == 1:
                img = block.get("image")
                if not img:
                    continue
                bb = block.get("bbox")
                try:
                    if bb and (bb[2] - bb[0]) > 1:
                        out.add_picture(io.BytesIO(img),
                                        width=Inches((bb[2] - bb[0]) / 72.0))
                    else:
                        out.add_picture(io.BytesIO(img))
                except Exception:
                    continue
        if progress:
            progress(i + 1, len(page_list))
    out.save(out_path)
    doc.close()
    return len(page_list)


def do_pdf_to_xlsx(src_path, out_path, pages=None, pw=None, progress=None):
    """PDF 转 Excel（.xlsx）：抽取可识别表格，每张表一个工作表；
    无表格的页按文本行写入一列。返回 (页数, 识别到的表格数)。"""
    from openpyxl import Workbook

    doc = open_pdf_doc(src_path, pw)
    page_list = list(range(doc.page_count)) if pages is None else list(pages)
    if not page_list:
        doc.close()
        raise ValueError("未选择任何页面。")
    wb = Workbook()
    wb.remove(wb.active)
    used = set()

    def new_sheet(name):
        base = name[:28]
        n, k = base, 2
        while n in used:
            n = f"{base}({k})"
            k += 1
        used.add(n)
        return wb.create_sheet(n)

    table_count = 0
    for pno in page_list:
        page = doc[pno]
        try:
            found = page.find_tables()
            tables = list(getattr(found, "tables", []) or [])
        except Exception:
            tables = []
        if tables:
            for ti, tab in enumerate(tables):
                table_count += 1
                ws = new_sheet(f"第{pno + 1}页-表{ti + 1}")
                try:
                    rows = tab.extract()
                except Exception:
                    rows = []
                for row in rows:
                    ws.append([cell_value(c) for c in row])
        else:
            ws = new_sheet(f"第{pno + 1}页")
            for line in page.get_text("text").splitlines():
                if line.strip():
                    ws.append([line])
        if progress:
            progress(pno + 1, len(page_list))
    if not wb.sheetnames:
        wb.create_sheet("无内容")
    wb.save(out_path)
    doc.close()
    return len(page_list), table_count


# ---------------- 书签 / 目录 ----------------

def get_toc_list(src_path, pw=None):
    """读取书签/目录，返回 [[层级, 标题, 目标页(1 基)], ...]。"""
    doc = open_pdf_doc(src_path, pw)
    toc = doc.get_toc()
    doc.close()
    return toc


def do_set_toc(src_path, out_path, toc, pw=None):
    """写入书签/目录。toc 为 [[层级(1..), 标题, 目标页(1 基)], ...]。"""
    doc = open_pdf_doc(src_path, pw)
    pages = doc.page_count
    norm = []
    for level, title, page in toc:
        lv = max(1, min(6, int(level)))
        t = str(title).strip()
        if not t:
            continue
        pg = max(1, min(pages, int(page)))
        norm.append([lv, t, pg])
    doc.set_toc(norm)
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return len(norm)


# ---------------- 表单填写 ----------------

def _fields_from_doc(doc):
    """从已打开的文档对象中读取所有可填写表单字段。"""
    out = []
    for page in doc:
        for w in (page.widgets() or []):
            try:
                choices = list(w.field_choices) if w.field_choices else []
            except Exception:
                choices = []
            out.append({
                "name": w.field_name or "",
                "type": w.field_type_string or "",
                "value": "" if w.field_value is None else str(w.field_value),
                "page": page.number + 1,
                "choices": choices,
            })
    return out


def list_form_fields(src_path, pw=None):
    """列出可填写的表单字段（名称/类型/当前值/所在页/可选值）。"""
    doc = open_pdf_doc(src_path, pw)
    fields = _fields_from_doc(doc)
    doc.close()
    return fields


def do_fill_form(src_path, out_path, values, pw=None):
    """填写表单字段并另存。values 为 {字段名: 值}。返回填写数量。"""
    doc = open_pdf_doc(src_path, pw)
    n = 0
    for page in doc:
        for w in (page.widgets() or []):
            fn = w.field_name
            if not fn or fn not in values:
                continue
            v = values[fn]
            if "CheckBox" in (w.field_type_string or ""):
                if isinstance(v, bool):
                    w.field_value = v
                else:
                    w.field_value = str(v).strip().lower() in (
                        "1", "true", "yes", "on", "是", "✓", "√")
            else:
                w.field_value = "" if v is None else str(v)
            try:
                w.update()
            except Exception:
                pass
            n += 1
    if n == 0:
        doc.close()
        raise ValueError("当前文档没有可填写的表单字段。")
    try:
        doc.need_appearances(True)
    except Exception:
        pass
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return n


# ---------------- 批量处理 ----------------

BATCH_OPS = ["添加文字水印", "添加页码", "压缩体积", "加密保护",
             "解密 / 去密码", "旋转页面"]


def _all_pages(src_path, pw=None):
    doc = open_pdf_doc(src_path, pw)
    n = doc.page_count
    doc.close()
    return list(range(n))


def do_batch(op, files, out_dir, params, progress=None):
    """批量处理：对多个 PDF 依次执行同一操作，各自另存到 out_dir（不覆盖）。
    返回 [(源文件, 输出文件, 错误信息), ...]，错误信息为空表示成功。"""
    os.makedirs(out_dir, exist_ok=True)
    results = []
    total = len(files)
    for i, src in enumerate(files):
        base = os.path.splitext(os.path.basename(src))[0]
        outp, err = "", ""
        try:
            if op == "添加文字水印":
                outp = unique_out_path(os.path.join(out_dir, f"{base}_水印.pdf"))
                do_watermark_text(
                    src, outp, None, params.get("text", "机密"),
                    params.get("fontsize", 48),
                    params.get("color", (0.75, 0.75, 0.75)),
                    params.get("angle", 45), params.get("opacity", 0.2),
                    params.get("layout", "平铺整页"))
            elif op == "添加页码":
                outp = unique_out_path(os.path.join(out_dir, f"{base}_加页码.pdf"))
                do_add_page_numbers(src, outp, params.get("position", "底部居中"),
                                    params.get("start", 1),
                                    params.get("fmt", "第 X 页 / 共 N 页"),
                                    params.get("fontsize", 10))
            elif op == "压缩体积":
                outp = unique_out_path(os.path.join(out_dir, f"{base}_压缩.pdf"))
                do_compress(src, outp, params.get("quality", 50))
            elif op == "加密保护":
                outp = unique_out_path(os.path.join(out_dir, f"{base}_加密.pdf"))
                do_encrypt(src, outp, params.get("password", "123456"))
            elif op == "解密 / 去密码":
                outp = unique_out_path(os.path.join(out_dir, f"{base}_无密码.pdf"))
                do_decrypt(src, outp, params.get("pw") or None)
            elif op == "旋转页面":
                outp = unique_out_path(os.path.join(out_dir, f"{base}_旋转.pdf"))
                do_rotate_pages(src, _all_pages(src), params.get("rotate", 90), outp)
            else:
                raise ValueError(f"不支持的操作：{op}")
        except Exception as e:
            err = str(e)
        results.append((src, outp, err))
        if progress:
            progress(i + 1, total)
    return results


# ============================================================
# V1.6 核心层：标注 / 遮盖 / 页眉页脚 / 提取图片
# ============================================================

ANNOT_TYPES = ["高亮", "下划线", "删除线", "波浪线", "便签", "文本框"]
ANNOT_COLORS = {
    "黄": (1.0, 0.90, 0.0),
    "红": (1.0, 0.20, 0.20),
    "蓝": (0.20, 0.50, 1.0),
    "绿": (0.20, 0.85, 0.30),
    "浅灰": (0.75, 0.75, 0.75),
}


def _rot_map(page_rot, w, h, u, v):
    """未旋转页面坐标 (u,v) → 旋转后（显示）坐标。w,h 为 mediabox 宽高。"""
    if page_rot == 90:
        return (h - v, u)
    if page_rot == 180:
        return (w - u, h - v)
    if page_rot == 270:
        return (v, w - u)
    return (u, v)


def _rot_unmap(page_rot, w, h, x, y):
    """旋转后（显示）坐标 (x,y) → 未旋转页面坐标。_rot_map 的逆运算。"""
    if page_rot == 90:
        return (y, h - x)
    if page_rot == 180:
        return (w - x, h - y)
    if page_rot == 270:
        return (w - y, x)
    return (x, y)


def page_geometry(page):
    """返回页面「显示坐标系」信息：(宽, 高, 旋转角, mediabox宽, mediabox高)。
    显示宽高已把 /Rotate 计入；未旋转坐标换算一律用 mediabox 宽高。
    实测（2026-10-17）：Page.rect 对 90/270 返回的是已旋转宽高，直接拿它配合
    insert_text 会错位——insert_text 用的是未旋转坐标系。"""
    mb = page.mediabox
    rot = (page.rotation or 0) % 360
    if rot in (90, 270):
        return mb.height, mb.width, rot, mb.width, mb.height
    return mb.width, mb.height, rot, mb.width, mb.height


def insert_text_display(page, x, y, text, fontname, fontsize, color=(0, 0, 0)):
    """在「显示坐标系」（旋转后、左上角为原点、文字水平正立）写字。
    实测（2026-10-17）：insert_text 落在未旋转坐标系；rotate=页面旋转角时，
    文字在显示中保持水平正立且自左向右排列。"""
    _dw, _dh, rot, mbw, mbh = page_geometry(page)
    u, v = _rot_unmap(rot, mbw, mbh, x, y)
    page.insert_text(fitz.Point(u, v), text, fontname=fontname,
                     fontsize=fontsize, color=color, rotate=rot)


def rect_display(page, rect):
    """未旋转矩形 → 显示坐标系矩形（用于搜索高亮、文字选择层）。"""
    _dw, _dh, rot, mbw, mbh = page_geometry(page)
    if not rot:
        return fitz.Rect(rect)
    pts = [_rot_map(rot, mbw, mbh, rect.x0, rect.y0),
           _rot_map(rot, mbw, mbh, rect.x1, rect.y0),
           _rot_map(rot, mbw, mbh, rect.x0, rect.y1),
           _rot_map(rot, mbw, mbh, rect.x1, rect.y1)]
    xs = [p[0] for p in pts]
    ys = [p[1] for p in pts]
    return fitz.Rect(min(xs), min(ys), max(xs), max(ys))


def image_physical_size(path):
    """按图片自身 DPI 换算物理尺寸（点）。缺 DPI 信息时按 96dpi（屏幕）。
    300dpi 的 A4 扫描图必须还原成 595x842pt，否则会被放大 3 倍。返回 (宽, 高)。"""
    from PIL import Image as PILImage
    try:
        with PILImage.open(path) as im:
            w, h = im.size
            dpi = im.info.get("dpi")
    except Exception:
        return None
    dx = dy = 96.0
    try:
        if dpi:
            dx, dy = float(dpi[0]), float(dpi[1])
    except Exception:
        dx = dy = 96.0
    if not (10 <= dx <= 2400):
        dx = 96.0
    if not (10 <= dy <= 2400):
        dy = 96.0
    return w * 72.0 / dx, h * 72.0 / dy


_NUM_CELL_RE = re.compile(r"^-?(0|[1-9]\d{0,9})(\.\d{1,8})?$")


def _is_cjk(s):
    """判断首字符是否中日韩表意文字（用于拼接选区文字时决定是否补空格）。"""
    if not s:
        return False
    c = s[0]
    return ("\u3000" <= c <= "\u9fff") or ("\uff00" <= c <= "\uffef")


def cell_value(v):
    """表格单元格转值：像数字的写成数字（Excel 里能求和），其余保持文本。
    限制整数位 ≤10 位，避免电话/单号/身份证被转成科学计数法。"""
    if v is None:
        return ""
    s = str(v).strip()
    if not s:
        return ""
    if _NUM_CELL_RE.match(s):
        return float(s) if "." in s else int(s)
    return s


def do_annotate(src_path, out_path, kind, areas=None, pages=None, keyword=None,
                text="", color=(1.0, 0.90, 0.0), fontsize=12, pw=None, progress=None):
    """添加标注/批注，另存新文件。返回添加数量。
    kind：高亮/下划线/删除线/波浪线/便签/文本框。
    areas：手动框选 [(页号0基, (x0,y0,x1,y1)), ...]（未旋转页面坐标）。
    keyword+pages：按关键词在各页中查找文字并自动标注（扫描件无效）。"""
    doc = open_pdf_doc(src_path, pw)
    targets = []  # [(页号0基, fitz.Rect)]
    if keyword:
        page_list = list(range(doc.page_count)) if pages is None else list(pages)
        for pno in page_list:
            for r in doc[pno].search_for(keyword):
                targets.append((pno, fitz.Rect(r)))
    else:
        for pno, box in (areas or []):
            targets.append((pno, fitz.Rect(box)))
    if not targets:
        doc.close()
        raise ValueError("没有找到可标注的区域。若为扫描件，关键词定位无效，请改用手动框选。")
    total = len(targets)
    for i, (pno, r) in enumerate(targets):
        page = doc[pno]
        if kind in ("高亮", "下划线", "删除线", "波浪线"):
            fn = {"高亮": page.add_highlight_annot, "下划线": page.add_underline_annot,
                  "删除线": page.add_strikeout_annot, "波浪线": page.add_squiggly_annot}[kind]
            a = fn(r)
            try:
                a.set_colors(stroke=color)
                a.update()
            except Exception:
                pass
        elif kind == "便签":
            a = page.add_text_annot(fitz.Point(r.x0, r.y0), text or "批注")
            try:
                a.set_colors(stroke=color)
                a.update()
            except Exception:
                pass
        else:  # 文本框
            box = fitz.Rect(r.x0 - 3, r.y0 - 3,
                            max(r.x1, r.x0) + 70, max(r.y1, r.y0) + 8)
            a = page.add_freetext_annot(box, text or "批注文字", fontsize=fontsize,
                                        fontname=FONT_CJK, text_color=(0, 0, 0),
                                        fill_color=(1.0, 1.0, 0.65))
            try:
                a.update()
            except Exception:
                pass
        if progress:
            progress(i + 1, total)
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return total


def do_ink(src_path, out_path, strokes, pw=None, progress=None):
    """画笔标记：把自由手绘笔迹按 PDF 点画进页面，另存为新文件。

    strokes：[{"page": 0基页号,
               "color": "#E11D48" 或 (r,g,b) 0..1,
               "width": 笔画粗细（点，缺省 3）,
               "points": [[x,y], ...]（未旋转页面坐标，PDF 点）}, ...]
    返回实际写入的笔画数。空笔画 / 越界页号一律跳过（不报错，方便前端批量提交）。
    """
    doc = open_pdf_doc(src_path, pw)
    n = apply_ink(doc, strokes, progress=progress)
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return n


def apply_ink(doc, strokes, progress=None):
    """把笔迹写进**已打开**的 doc（内容级墨迹，任何阅读器都能看见，不需交互式批注）。

    坐标用未旋转页面坐标（与 insert_text / search_for 同一套口径），
    旋转页由调用方换算好再传进来。返回写入的笔画数。

    每笔可带 shape：
      "free"（缺省）自由手绘折线；"line" 直线（首尾两点）；
      "rect" 矩形（首尾两点为对角）；"ellipse" 椭圆（首尾两点为外接矩形）。
    未知 shape 一律按 free 处理。
    """
    total = 0
    done = 0
    strokes = strokes or []
    n_pages = doc.page_count
    for s in strokes:
        pno = int(s.get("page", 0) or 0)
        pts = [(float(x), float(y)) for x, y in (s.get("points") or [])]
        if pno < 0 or pno >= n_pages or not pts:
            continue
        kind = str(s.get("shape") or "free").lower()
        if kind not in ("free", "line", "rect", "ellipse"):
            kind = "free"
        if kind != "free" and len(pts) >= 2:
            p0, p1 = pts[0], pts[-1]
            box = fitz.Rect(min(p0[0], p1[0]), min(p0[1], p1[1]),
                            max(p0[0], p1[0]), max(p0[1], p1[1]))
        try:
            color = hex_to_rgb(s.get("color") or INK_PRESETS["红"])
        except ValueError:
            color = hex_to_rgb(INK_PRESETS["红"])
        width = max(0.5, float(s.get("width") or 3))
        page = doc[pno]
        shape = page.new_shape()
        if kind == "line":
            shape.draw_line(fitz.Point(*pts[0]), fitz.Point(*pts[-1]))
        elif kind == "rect":
            shape.draw_rect(box)
        elif kind == "ellipse":
            shape.draw_oval(box)
        else:
            shape.draw_polyline([fitz.Point(x, y) for x, y in pts])
        # lineCap/lineJoin 取圆角，单点也能画出一个圆点（手写落笔感）
        shape.finish(color=color, width=width, lineCap=1, lineJoin=1)
        shape.commit()
        total += 1
        done += 1
        if progress:
            progress(done, len(strokes))
    return total


def do_save_as(src_path, out_path, pw=None, strokes=None, progress=None):
    """另存为新文件：可顺带写画笔笔迹，原文件保持不动。返回 (路径, 笔迹数)。"""
    if out_path and os.path.abspath(out_path) == os.path.abspath(src_path):
        raise ValueError("输出路径不能等于输入文件（另存请换个名字）")
    doc = open_pdf_doc(src_path, pw)
    n = apply_ink(doc, strokes, progress=progress) if strokes else 0
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return out_path, n


def save_doc_overwrite(doc, src_path):
    """把已打开的 doc 写回它自己的路径（临时文件 + os.replace）。

    Windows 上不能 doc.save(自身路径)（文件被占用，PyMuPDF 直接抛错），
    也不能在 doc 还开着时 replace（句柄占着原 inode）——所以顺序固定为：
    save 到同目录临时文件 → close doc → os.replace。调用后 doc 已关闭，
    主界面需重新 open_file 才会拿到新内容。
    """
    src_path = os.path.abspath(src_path)
    tmp = src_path + ".3qsave.tmp"
    kw = dict(garbage=3, deflate=True)
    try:
        if doc.is_encrypted:
            # 保持原密码与权限，别把加密文档保存成明文
            kw["encryption"] = fitz.PDF_ENCRYPT_KEEP
    except Exception:
        pass
    doc.save(tmp, **kw)
    doc.close()
    last = None
    for _ in range(5):
        try:
            os.replace(tmp, src_path)
            return src_path
        except OSError as e:      # 缩略图线程等还握着旧文件：退避重试
            last = e
            time.sleep(0.2)
    try:
        os.remove(tmp)
    except OSError:
        pass
    raise last if last else RuntimeError("写回原文件失败")


def do_save_inplace(src_path, pw=None, strokes=None, progress=None):
    """覆盖保存（CLI / MCP 用）：可选顺带写入画笔笔迹，然后写回原路径。

    这是「直接保存」能力的核心：不再强制用户另存新文件。返回写回路径。
    """
    doc = open_pdf_doc(src_path, pw)
    n = apply_ink(doc, strokes, progress=progress) if strokes else 0
    path = save_doc_overwrite(doc, src_path)
    return path, n


def do_redact(src_path, out_path, areas=None, pages=None, keyword=None,
              fill=(0, 0, 0), pw=None, progress=None):
    """涂黑遮盖：框选或按关键词遮盖区域，并抹除其下文字（不可复制）。返回遮盖数量。
    areas：[(页号0基, (x0,y0,x1,y1)), ...]；keyword+pages：按关键词定位。"""
    doc = open_pdf_doc(src_path, pw)
    per_page = {}
    if keyword:
        page_list = list(range(doc.page_count)) if pages is None else list(pages)
        for pno in page_list:
            hits = doc[pno].search_for(keyword)
            if hits:
                per_page.setdefault(pno, []).extend(hits)
    else:
        for pno, box in (areas or []):
            per_page.setdefault(pno, []).append(fitz.Rect(box))
    if not per_page:
        doc.close()
        raise ValueError("没有找到可遮盖的区域。若为扫描件，关键词定位无效，请改用手动框选。")
    total = sum(len(v) for v in per_page.values())
    count = 0
    for pno, rects in per_page.items():
        page = doc[pno]
        for r in rects:
            rr = fitz.Rect(r.x0 - 0.5, r.y0 - 0.5, r.x1 + 0.5, r.y1 + 0.5)
            page.add_redact_annot(rr, fill=fill)
            count += 1
        page.apply_redactions()
        if progress:
            progress(count, total)
    # 遮盖 ≠ 净化：只抹页面内容时，文档属性/书签/XMP/隐藏文字里仍能捞回原文
    # （实测：不清时原始字节里仍搜得到被遮盖的敏感串）。这里连文档级残留一起清掉。
    try:
        doc.scrub(attached_files=True, clean_pages=True, embedded_files=True,
                  hidden_text=True, javascript=True, metadata=True,
                  redactions=True, remove_links=True, reset_fields=True,
                  reset_responses=True, thumbnails=True, xml_metadata=True)
    except Exception:
        try:
            doc.set_metadata({})
            doc.del_xml_metadata()
        except Exception:
            pass
    try:
        if doc.get_toc():
            doc.set_toc([])  # 书签标题常含原文，一并清空
    except Exception:
        pass
    doc.save(out_path, garbage=4, deflate=True, clean=True)
    doc.close()
    return count


def do_header_footer(src_path, out_path, cfg, pw=None):
    """添加页眉/页脚，另存新文件。cfg = {header, footer, size, margin_mm,
    align, pages, start, color}。文本中 {page} 为当前页号，{total} 为总页数。"""
    doc = open_pdf_doc(src_path, pw)
    pages = cfg.get("pages") or list(range(doc.page_count))
    header = cfg.get("header", "") or ""
    footer = cfg.get("footer", "") or ""
    size = max(6, min(72, int(cfg.get("size", 10))))
    margin = max(3, float(cfg.get("margin_mm", 10))) * MM
    align = cfg.get("align", "居中")
    color = cfg.get("color", (0, 0, 0))
    start = int(cfg.get("start", 1))
    n = doc.page_count
    font = fitz.Font(FONT_CJK)
    for pno in pages:
        page = doc[pno]
        dw, dh = page_geometry(page)[:2]
        num = pno + start
        for which, txt in (("header", header), ("footer", footer)):
            t = (txt or "").strip()
            if not t:
                continue
            s = t.replace("{page}", str(num)).replace("{total}", str(n))
            fs = _fit_fontsize(font, s, size, dw * 0.9)
            w = text_width(s, fs)
            if align == "左侧":
                x = margin
            elif align == "右侧":
                x = dw - margin - w
            else:
                x = (dw - w) / 2
            y = margin + fs if which == "header" else dh - margin
            insert_text_display(page, x, y, s, FONT_CJK, fs, color)
    doc.save(out_path, garbage=3, deflate=True)
    doc.close()
    return len(pages)


def do_extract_images(src_path, out_dir, pw=None, want_fmt="auto", min_px=16, progress=None):
    """提取文档内嵌图片：每页引用到的图片 XObject 导出为文件。返回输出路径列表。
    want_fmt：auto 保持原格式；png/jpg 统一转换（jpg 无透明通道）。"""
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.splitext(os.path.basename(src_path))[0]
    doc = open_pdf_doc(src_path, pw)
    outs = []
    seen = set()
    for pno in range(doc.page_count):
        page = doc[pno]
        for img in page.get_images(full=True):
            xref = img[0]
            if xref in seen:
                continue
            try:
                info = doc.extract_image(xref)
            except Exception:
                continue
            data = info.get("image") or b""
            if not data:
                continue
            if min_px and (info.get("width", 0) < min_px or info.get("height", 0) < min_px):
                continue
            ext = info.get("ext", "png")
            if want_fmt in ("png", "jpg"):
                from PIL import Image as PILImage
                try:
                    im = PILImage.open(io.BytesIO(data))
                    buf = io.BytesIO()
                    if want_fmt == "jpg":
                        if im.mode not in ("RGB", "L"):
                            im = im.convert("RGB")
                        im.save(buf, "JPEG", quality=90)
                    else:
                        im.save(buf, "PNG")
                    data = buf.getvalue()
                    ext = want_fmt
                except Exception:
                    pass  # 转换失败则保持原格式
            p = unique_out_path(
                os.path.join(out_dir, f"{base}_第{pno + 1}页_图{len(outs) + 1}.{ext}"))
            try:
                with open(p, "wb") as f:
                    f.write(data)
            except OSError:
                continue
            seen.add(xref)
            outs.append(p)
        if progress:
            progress(pno + 1, doc.page_count)
    doc.close()
    return outs


# ============================================================
# FR-06 查看与打印（核心层）
# ============================================================

def page_pixmap_png(doc, pno, zoom):
    """渲染指定页为 PNG 字节流（供界面显示）。返回 (png字节, 宽px, 高px)。"""
    pix = doc[pno].get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    return pix.tobytes("png"), pix.width, pix.height


def list_printers():
    """枚举本机打印机。返回 (打印机名列表, 默认打印机名)。"""
    import win32print
    names = [p[2] for p in win32print.EnumPrinters(
        win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS)]
    try:
        default = win32print.GetDefaultPrinter()
    except Exception:
        default = names[0] if names else ""
    return names, default


def do_print(src_path, printer_name, page_indices, copies=1, output_file=None,
             progress=None, pw=None):
    """FR-06 打印：将指定页渲染后送往打印机，按纸张可打印区域等比缩放居中。
    output_file 仅供自动化测试使用（如 Microsoft Print to PDF 输出到文件）。"""
    import win32ui
    import win32con
    from PIL import Image, ImageWin

    doc = open_pdf_doc(src_path, pw)
    hDC = win32ui.CreateDC()
    hDC.CreatePrinterDC(printer_name)
    try:
        dpi = hDC.GetDeviceCaps(win32con.LOGPIXELSX)
        pw_ = hDC.GetDeviceCaps(win32con.HORZRES)
        ph = hDC.GetDeviceCaps(win32con.VERTRES)
        render_dpi = max(150, min(dpi, 300))
        copies = max(1, int(copies))
        total = len(page_indices) * copies
        done = 0
        hDC.StartDoc("3Q PDF开源免费工具打印", output_file)
        try:
            for _ in range(copies):
                for pno in page_indices:
                    page = doc[pno]
                    r = page.rect
                    pix = page.get_pixmap(dpi=render_dpi, alpha=False)
                    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    page_w_dev = r.width / 72.0 * dpi
                    page_h_dev = r.height / 72.0 * dpi
                    scale = min(pw_ / page_w_dev, ph / page_h_dev)
                    dw = max(1, int(page_w_dev * scale))
                    dh = max(1, int(page_h_dev * scale))
                    x, y = (pw_ - dw) // 2, (ph - dh) // 2
                    hDC.StartPage()
                    ImageWin.Dib(img).draw(hDC.GetHandleAttrib(),
                                           (x, y, x + dw, y + dh))
                    hDC.EndPage()
                    done += 1
                    if progress:
                        progress(done, total)
            hDC.EndDoc()
        except Exception:
            try:
                hDC.AbortDoc()
            except Exception:
                pass
            raise
    finally:
        hDC.DeleteDC()
        doc.close()
    return total
